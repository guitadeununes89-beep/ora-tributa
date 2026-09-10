"""End-to-end batch classification against real deployed data (Etapa 23, 24).

Requires the same deployment `test_evaluate_many.py`/
`test_catalog_discovery_unified_integration.py` already require: the 5 real
rules and the NCM/NBS catalogs published to `DATABASE_URL` (see
docs/DEVELOPMENT_ACCESS.md). Proves the whole pipeline end to end without a
second motor: NCM lookup -> discovery -> ruleset resolution -> the real,
unmodified `EvaluationService.evaluate_many()` -> a genuine persisted
`EvaluationRecord`.

Etapa 24 (ADR-0028) made `process`/`reprocess` asynchronous: they now only
enqueue `run_batch_job` via `BackgroundTasks` and return immediately. Tests
call the route (to prove it returns right away with `PROCESSING`), then
invoke `run_batch_job` directly - the same pattern FastAPI's own docs
recommend for testing background tasks - to actually run the work, then
`get_batch` for the final assertions.
"""

from __future__ import annotations

import asyncio
import io
import os
from typing import Any

import pytest
from fastapi import BackgroundTasks
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from starlette.datastructures import UploadFile
from tributaria_api.api.routes.batch_classification import (
    BatchDetailResponse,
    export,
    get_batch,
    process,
    reprocess,
    run_batch_job,
    upload,
)
from tributaria_api.application.errors import NotFoundError
from tributaria_api.application.security import AuthContext, Role
from tributaria_api.infrastructure.database.batch_repository import SqlAlchemyBatchRepository
from tributaria_api.infrastructure.database.nbs_repository import SqlAlchemyNbsRepository
from tributaria_api.infrastructure.database.ncm_repository import SqlAlchemyNcmRepository
from tributaria_api.infrastructure.database.product_repository import ProductRepository
from tributaria_api.infrastructure.database.repositories import SqlAlchemyGovernanceRepository
from tributaria_api.infrastructure.database.taxonomy_repository import (
    SqlAlchemyTaxonomyRepository,
)

pytestmark = pytest.mark.skipif(
    os.getenv("POSTGRES_TESTS") != "1",
    reason="requires migrated PostgreSQL with real rules and catalogs deployed",
)

_ORGANIZATION_ID = "dev-governance-org"
_ANALYST_USER_ID = "dev-analyst"


def _context(organization_id: str = _ORGANIZATION_ID) -> AuthContext:
    return AuthContext(
        session_id="session",
        user_id=_ANALYST_USER_ID,
        email="analyst@example.invalid",
        display_name="Synthetic analyst",
        organization_id=organization_id,
        organization_name="Synthetic organization",
        membership_id="membership",
        role=Role.ANALYST,
    )


def _rt_0005_eligible_facts() -> dict[str, str]:
    return {
        "product.kind": "MEDICINE",
        "product.anvisa_registration_status": "REGISTERED",
        "buyer.health_entity_status": "HEALTH_ENTITY",
        "buyer.ibs_cbs_immunity_status": "IMMUNE",
        "operation.effective_buyer_status": "CONFIRMED",
        "buyer.cebas_status": "VALID",
        "buyer.sus_service_requirement_status": "SATISFIED",
    }


def _xlsx_bytes(rows: list[list[object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


async def _read_streaming_response(response: object) -> bytes:
    chunks = []
    async for chunk in response.body_iterator:  # type: ignore[attr-defined]
        chunks.append(chunk if isinstance(chunk, bytes) else chunk.encode())
    return b"".join(chunks)


def _repositories(session: Session) -> dict[str, Any]:
    return {
        "batches": SqlAlchemyBatchRepository(session),
        "repository": SqlAlchemyGovernanceRepository(session),
        "ncm_repository": SqlAlchemyNcmRepository(session),
        "nbs_repository": SqlAlchemyNbsRepository(session),
        "products": ProductRepository(session),
        "taxonomy": SqlAlchemyTaxonomyRepository(session),
    }


def _skip_if_prerequisites_missing(repositories: dict[str, Any]) -> None:
    governance = repositories["repository"]
    try:
        governance.get_ruleset("IBSCBS-PILOT-0004-001")
        governance.get_ruleset("IBSCBS-PILOT-0005-001")
    except NotFoundError:
        pytest.skip("RT-IBSCBS-0004/0005 not deployed - run the real rule deploy CLIs first")
    try:
        repositories["ncm_repository"].get_code(_ORGANIZATION_ID, "30019010")
    except NotFoundError:
        pytest.skip("NCM catalog not deployed - run governed_ncm_nbs_load_cli first")


def _process_and_wait(
    batch_id: str, *, correlation_id: str, repositories: dict[str, Any]
) -> BatchDetailResponse:
    """Call `process`, then run its enqueued job synchronously, then fetch.

    Mirrors exactly what Starlette does for a real request: run the route
    (which only enqueues), then run the background task, then let a
    separate `GET` observe the result - proving the async wiring actually
    reaches the same end state the old synchronous flow did.
    """
    background_tasks = BackgroundTasks()
    summary = process(
        batch_id,
        _context(),
        correlation_id,
        background_tasks,
        repositories["batches"],
        repositories["ncm_repository"],
        repositories["nbs_repository"],
        repositories["taxonomy"],
    )
    assert summary.status == "PROCESSING"
    run_batch_job(batch_id, organization_id=_ORGANIZATION_ID, correlation_id=correlation_id)
    return get_batch(batch_id, _context(), repositories["batches"], repositories["repository"])


def _reprocess_and_wait(
    batch_id: str, *, correlation_id: str, repositories: dict[str, Any]
) -> BatchDetailResponse:
    background_tasks = BackgroundTasks()
    summary = reprocess(
        batch_id,
        _context(),
        correlation_id,
        background_tasks,
        repositories["batches"],
        repositories["ncm_repository"],
        repositories["nbs_repository"],
        repositories["taxonomy"],
    )
    assert summary.status == "PROCESSING"
    run_batch_job(summary.batch_id, organization_id=_ORGANIZATION_ID, correlation_id=correlation_id)
    return get_batch(
        summary.batch_id, _context(), repositories["batches"], repositories["repository"]
    )


def test_batch_upload_and_processing_produces_a_conclusive_row() -> None:
    engine = create_engine(os.environ["DATABASE_URL"])
    with Session(engine) as session:
        repositories = _repositories(session)
        _skip_if_prerequisites_missing(repositories)

        content = _xlsx_bytes(
            [
                [
                    "Código Interno",
                    "Descrição",
                    "NCM",
                    "Data da Operação",
                    *_rt_0005_eligible_facts().keys(),
                ],
                [
                    "SKU-BATCH-1",
                    "Heparina e seus sais",
                    "30019010",
                    "2026-01-05",
                    *_rt_0005_eligible_facts().values(),
                ],
            ]
        )
        upload_file = UploadFile(file=io.BytesIO(content), filename="lote.xlsx")
        upload_response = asyncio.run(
            upload(_context(), repositories["batches"], upload_file)  # type: ignore[arg-type]
        )
        assert upload_response.batch.row_count == 1
        assert upload_response.batch.status == "RECEIVED"

        detail = _process_and_wait(
            upload_response.batch.batch_id,
            correlation_id="test-correlation-batch-1",
            repositories=repositories,
        )

        assert detail.batch.status == "COMPLETED"
        assert detail.batch.processed_count == 1
        assert detail.batch.error_count == 0
        row = detail.rows[0]
        assert row.processing_status == "PROCESSED"
        assert row.classification_status == "CONCLUSIVO"
        assert row.cclasstrib == "200010"  # RT-IBSCBS-0005 (supported); 0004 is SCOPE_UNCONFIRMED
        assert row.evaluation_id is not None

        stored = repositories["repository"].get_evaluation(row.evaluation_id)
        assert sorted(stored["composed_ruleset_ids"]) == [
            "IBSCBS-PILOT-0004-001",
            "IBSCBS-PILOT-0005-001",
        ]


def test_ncm_outside_chapter_30_has_no_governed_coverage() -> None:
    engine = create_engine(os.environ["DATABASE_URL"])
    with Session(engine) as session:
        repositories = _repositories(session)
        _skip_if_prerequisites_missing(repositories)

        content = _xlsx_bytes(
            [
                ["NCM", "Data da Operação"],
                ["01012100", "2026-01-05"],
            ]
        )
        upload_file = UploadFile(file=io.BytesIO(content), filename="lote.xlsx")
        upload_response = asyncio.run(
            upload(_context(), repositories["batches"], upload_file)  # type: ignore[arg-type]
        )

        detail = _process_and_wait(
            upload_response.batch.batch_id,
            correlation_id="test-correlation-batch-2",
            repositories=repositories,
        )

        row = detail.rows[0]
        assert row.processing_status == "PROCESSED"
        assert row.classification_status == "SEM_COBERTURA_NORMATIVA"
        assert row.evaluation_id is None


def test_ncm_not_in_published_catalog_is_an_isolated_row_error() -> None:
    engine = create_engine(os.environ["DATABASE_URL"])
    with Session(engine) as session:
        repositories = _repositories(session)
        _skip_if_prerequisites_missing(repositories)

        content = _xlsx_bytes(
            [
                ["NCM", "Data da Operação"],
                ["99999999", "2026-01-05"],
                ["30019010", "2026-01-05"],
            ]
        )
        upload_file = UploadFile(file=io.BytesIO(content), filename="lote.xlsx")
        upload_response = asyncio.run(
            upload(_context(), repositories["batches"], upload_file)  # type: ignore[arg-type]
        )

        detail = _process_and_wait(
            upload_response.batch.batch_id,
            correlation_id="test-correlation-batch-3",
            repositories=repositories,
        )

        # The bad row does not interrupt the good one that follows it.
        assert detail.batch.error_count == 1
        assert detail.batch.processed_count == 1
        bad_row, good_row = detail.rows
        assert bad_row.processing_status == "ERROR"
        assert good_row.processing_status == "PROCESSED"


def test_batch_isolation_between_organizations() -> None:
    engine = create_engine(os.environ["DATABASE_URL"])
    with Session(engine) as session:
        repositories = _repositories(session)
        _skip_if_prerequisites_missing(repositories)
        batches = repositories["batches"]

        content = _xlsx_bytes([["NCM", "Data da Operação"], ["30019010", "2026-01-05"]])
        upload_file = UploadFile(file=io.BytesIO(content), filename="lote.xlsx")
        upload_response = asyncio.run(upload(_context(), batches, upload_file))  # type: ignore[arg-type]

        with pytest.raises(NotFoundError):
            batches.get_batch("dev-org", upload_response.batch.batch_id)
        # The owning organization can still see it.
        batches.get_batch(_ORGANIZATION_ID, upload_response.batch.batch_id)


def test_reprocessing_a_batch_reproduces_the_same_classification() -> None:
    engine = create_engine(os.environ["DATABASE_URL"])
    with Session(engine) as session:
        repositories = _repositories(session)
        _skip_if_prerequisites_missing(repositories)

        content = _xlsx_bytes(
            [
                ["NCM", "Data da Operação", *_rt_0005_eligible_facts().keys()],
                ["30019010", "2026-01-05", *_rt_0005_eligible_facts().values()],
            ]
        )
        upload_file = UploadFile(file=io.BytesIO(content), filename="lote.xlsx")
        first_upload = asyncio.run(
            upload(_context(), repositories["batches"], upload_file)  # type: ignore[arg-type]
        )
        first_detail = _process_and_wait(
            first_upload.batch.batch_id,
            correlation_id="test-correlation-reprocess-1",
            repositories=repositories,
        )

        second_detail = _reprocess_and_wait(
            first_upload.batch.batch_id,
            correlation_id="test-correlation-reprocess-2",
            repositories=repositories,
        )

        first_row, second_row = first_detail.rows[0], second_detail.rows[0]
        assert first_row.classification_status == second_row.classification_status
        assert first_row.cclasstrib == second_row.cclasstrib
        assert first_row.evaluation_id != second_row.evaluation_id


def test_export_produces_a_readable_xlsx_with_original_and_result_columns() -> None:
    engine = create_engine(os.environ["DATABASE_URL"])
    with Session(engine) as session:
        repositories = _repositories(session)
        _skip_if_prerequisites_missing(repositories)

        content = _xlsx_bytes(
            [
                ["Código Interno", "NCM", "Data da Operação"],
                ["SKU-EXPORT-1", "30019010", "2026-01-05"],
            ]
        )
        upload_file = UploadFile(file=io.BytesIO(content), filename="lote.xlsx")
        upload_response = asyncio.run(
            upload(_context(), repositories["batches"], upload_file)  # type: ignore[arg-type]
        )
        _process_and_wait(
            upload_response.batch.batch_id,
            correlation_id="test-correlation-export",
            repositories=repositories,
        )

        response = export(
            upload_response.batch.batch_id,
            _context(),
            repositories["batches"],
            repositories["repository"],
        )
        body = asyncio.run(_read_streaming_response(response))

        workbook = load_workbook(io.BytesIO(body))
        sheet = workbook.active
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        assert "Código Interno" in headers
        assert "status" in headers
        data_row = [cell.value for cell in next(sheet.iter_rows(min_row=2, max_row=2))]
        assert "SKU-EXPORT-1" in data_row


def test_process_returns_immediately_without_waiting_for_row_evaluation() -> None:
    """Etapa 24 (ADR-0028): the HTTP response must never block on row work.

    Calling `process` alone (without running `run_batch_job`) must leave
    every row untouched (`PENDING`) - proof the response really came back
    before any evaluation happened, not just that it happens to be fast.
    """
    engine = create_engine(os.environ["DATABASE_URL"])
    with Session(engine) as session:
        repositories = _repositories(session)
        _skip_if_prerequisites_missing(repositories)

        content = _xlsx_bytes([["NCM", "Data da Operação"], ["30019010", "2026-01-05"]])
        upload_file = UploadFile(file=io.BytesIO(content), filename="lote.xlsx")
        upload_response = asyncio.run(
            upload(_context(), repositories["batches"], upload_file)  # type: ignore[arg-type]
        )

        background_tasks = BackgroundTasks()
        summary = process(
            upload_response.batch.batch_id,
            _context(),
            "test-correlation-no-wait",
            background_tasks,
            repositories["batches"],
            repositories["ncm_repository"],
            repositories["nbs_repository"],
            repositories["taxonomy"],
        )

        assert summary.status == "PROCESSING"
        assert summary.processed_count == 0
        assert len(background_tasks.tasks) == 1
        enqueued = background_tasks.tasks[0]
        assert enqueued.func is run_batch_job
        assert enqueued.kwargs["batch_id"] == upload_response.batch.batch_id

        row = repositories["batches"].list_rows(upload_response.batch.batch_id)[0]
        assert row.processing_status == "PENDING"


def test_unhandled_exception_in_the_job_marks_the_batch_failed(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    engine = create_engine(os.environ["DATABASE_URL"])
    with Session(engine) as session:
        repositories = _repositories(session)
        _skip_if_prerequisites_missing(repositories)

        content = _xlsx_bytes([["NCM", "Data da Operação"], ["30019010", "2026-01-05"]])
        upload_file = UploadFile(file=io.BytesIO(content), filename="lote.xlsx")
        upload_response = asyncio.run(
            upload(_context(), repositories["batches"], upload_file)  # type: ignore[arg-type]
        )

        background_tasks = BackgroundTasks()
        process(
            upload_response.batch.batch_id,
            _context(),
            "test-correlation-failure",
            background_tasks,
            repositories["batches"],
            repositories["ncm_repository"],
            repositories["nbs_repository"],
            repositories["taxonomy"],
        )

        def _raise(*args: Any, **kwargs: Any) -> Any:
            raise RuntimeError("simulated unexpected failure mid-batch")

        monkeypatch.setattr(
            "tributaria_api.api.routes.batch_classification.process_batch_row", _raise
        )

        run_batch_job(
            upload_response.batch.batch_id,
            organization_id=_ORGANIZATION_ID,
            correlation_id="test-correlation-failure",
        )

    # New session: confirm the failure was committed, not just held in the
    # (rolled-back) session the job used.
    with Session(engine) as session:
        repositories = _repositories(session)
        batch = repositories["batches"].get_batch(_ORGANIZATION_ID, upload_response.batch.batch_id)
        assert batch.status == "FAILED"
        assert batch.completed_at is not None


def test_completed_batch_rows_are_immutable_even_to_a_late_duplicate_write() -> None:
    """The trigger from Etapa 23 must still hold once a job actually finishes -
    a duplicate/late invocation of `run_batch_job` for an already-`COMPLETED`
    batch must never be able to silently rewrite a stored result."""
    engine = create_engine(os.environ["DATABASE_URL"])
    with Session(engine) as session:
        repositories = _repositories(session)
        _skip_if_prerequisites_missing(repositories)

        content = _xlsx_bytes([["NCM", "Data da Operação"], ["30019010", "2026-01-05"]])
        upload_file = UploadFile(file=io.BytesIO(content), filename="lote.xlsx")
        upload_response = asyncio.run(
            upload(_context(), repositories["batches"], upload_file)  # type: ignore[arg-type]
        )
        detail = _process_and_wait(
            upload_response.batch.batch_id,
            correlation_id="test-correlation-immutable",
            repositories=repositories,
        )
        assert detail.batch.status == "COMPLETED"

    with Session(engine) as session:
        batches = SqlAlchemyBatchRepository(session)
        # Values must differ from what is already stored (1/0) - otherwise
        # SQLAlchemy's dirty-tracking sees no net change and skips emitting
        # an UPDATE at all, so the trigger would never even get a chance to
        # fire.
        with pytest.raises(Exception, match="immutable"):
            batches.update_progress(
                upload_response.batch.batch_id, processed_count=999, error_count=0
            )
