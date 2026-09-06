from __future__ import annotations

import hashlib
import json
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from lxml import etree
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

CST_HEADERS = (
    "CST-IBS/CBS",
    "Descrição CST-IBS/CBS",
    "ind_gIBSCBS",
    "ind_gIBSCBSMono",
    "ind_gRed",
    "ind_gDif",
    "ind_gTransfCred",
    "ind_ gCredPresIBSZFM",
    "ind_gAjusteCompet",
    "ind_RedutorBC",
)

CCLASSTRIB_HEADERS = (
    "CST-IBS/CBS",
    "Descrição CST-IBS/CBS",
    "cClassTrib",
    "Nome cClassTrib",
    "Descrição cClassTrib",
    "LC Redação",
    "LC 214/25",
    "Regulamento CBS",
    "Regulamento IBS",
    "Tipo de Alíquota",
    "pRedIBS",
    "pRedCBS",
    "ind_gTribRegular",
    "ind_gCredPresOper",
    "ind_gMonoPadrao",
    "ind_gMonoReten",
    "ind_gMonoRet",
    "ind_gMonoDif",
    "ind_gpBioDiferenca",
    "ind_gEstornoCred",
    "tpRBSN",
    "dIniVig",
    "dFimVig",
    "DataAtualização",
    "indNFeABI",
    "indNFe",
    "indNFCe",
    "indCTe",
    "indCTeOS",
    "indBPe",
    "indBPeTA",
    "indBPeTM",
    "indNF3e",
    "indNFSe",
    "indNFSe Via",
    "indNFCom",
    "indNFAg",
    "indNFGas",
    "indDERE",
    "indDIR",
    "indDUIMP",
    "ANEXO",
    "Link",
)

CCLASSTRIB_HEADERS_LEGACY = tuple(
    header
    for header in CCLASSTRIB_HEADERS
    if header
    not in {
        "Regulamento CBS",
        "Regulamento IBS",
        "ind_gpBioDiferenca",
        "tpRBSN",
        "indDIR",
        "indDUIMP",
    }
)

MAX_COMPRESSED_BYTES = 20_000_000
MAX_UNCOMPRESSED_BYTES = 100_000_000
MAX_ZIP_ENTRIES = 500


class CatalogImportError(ValueError):
    """The artifact cannot advance because its structure or content is unsafe."""


@dataclass(frozen=True, slots=True)
class ImportIssue:
    code: str
    message: str
    sheet: str | None = None
    row_number: int | None = None
    field: str | None = None


@dataclass(frozen=True, slots=True)
class StagingRow:
    sheet: str
    row_number: int
    raw: dict[str, Any]
    errors: tuple[ImportIssue, ...] = ()


@dataclass(frozen=True, slots=True)
class ParsedCatalog:
    artifact_hash: str
    normalized_hash: str
    schema_signature: str
    csts: tuple[dict[str, Any], ...]
    classifications: tuple[dict[str, Any], ...]
    staging_rows: tuple[StagingRow, ...]
    issues: tuple[ImportIssue, ...] = ()
    sheet_names: tuple[str, ...] = ()

    @property
    def valid(self) -> bool:
        return not self.issues and all(not row.errors for row in self.staging_rows)

    @property
    def report(self) -> dict[str, Any]:
        return {
            "valid": self.valid,
            "cst_count": len(self.csts),
            "cclasstrib_count": len(self.classifications),
            "staging_row_count": len(self.staging_rows),
            "issue_count": len(self.issues) + sum(len(row.errors) for row in self.staging_rows),
            "schema_signature": self.schema_signature,
            "sheets": list(self.sheet_names),
            "issues": [
                {
                    "code": issue.code,
                    "message": issue.message,
                    "sheet": issue.sheet,
                    "row_number": issue.row_number,
                    "field": issue.field,
                }
                for issue in self.issues
            ],
        }


def import_official_catalog(path: str | Path) -> ParsedCatalog:
    artifact = Path(path)
    _validate_container(artifact)
    artifact_bytes = artifact.read_bytes()
    artifact_hash = hashlib.sha256(artifact_bytes).hexdigest()
    workbook = load_workbook(artifact, read_only=True, data_only=True, keep_links=False)
    try:
        sheet_headers = {
            sheet.title: _trim_headers(
                tuple(_text(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1)))
            )
            for sheet in workbook.worksheets
        }
        cst_sheet = _sheet_for_headers(sheet_headers, CST_HEADERS)
        class_sheet, class_headers = _sheet_for_header_variants(
            sheet_headers, (CCLASSTRIB_HEADERS, CCLASSTRIB_HEADERS_LEGACY)
        )
        schema_signature = _schema_signature(sheet_headers)
        csts, cst_staging, cst_issues = _parse_csts(workbook[cst_sheet])
        classifications, class_staging, class_issues = _parse_classifications(
            workbook[class_sheet],
            {item["code"] for item in csts},
            workbook.epoch,
            class_headers,
        )
    finally:
        workbook.close()
    issues = tuple(cst_issues + class_issues)
    canonical = {
        "schema_signature": schema_signature,
        "csts": sorted(csts, key=lambda item: item["code"]),
        "classifications": sorted(classifications, key=lambda item: item["code"]),
    }
    normalized_hash = hashlib.sha256(_canonical_json(canonical).encode()).hexdigest()
    return ParsedCatalog(
        artifact_hash=artifact_hash,
        normalized_hash=normalized_hash,
        schema_signature=schema_signature,
        csts=tuple(csts),
        classifications=tuple(classifications),
        staging_rows=tuple(cst_staging + class_staging),
        issues=issues,
        sheet_names=tuple(sheet_headers),
    )


def _validate_container(path: Path) -> None:
    if not path.is_file() or path.stat().st_size > MAX_COMPRESSED_BYTES:
        raise CatalogImportError("Artifact is missing or exceeds the compressed size limit")
    if not zipfile.is_zipfile(path):
        raise CatalogImportError("Artifact is not an XLSX ZIP container")
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > MAX_ZIP_ENTRIES:
            raise CatalogImportError("XLSX contains too many ZIP entries")
        total = sum(entry.file_size for entry in entries)
        if total > MAX_UNCOMPRESSED_BYTES:
            raise CatalogImportError("XLSX exceeds the uncompressed size limit")
        unsafe = [entry.filename for entry in entries if ".." in Path(entry.filename).parts]
        if unsafe:
            raise CatalogImportError("XLSX contains an unsafe ZIP path")
        relationship_files = [entry for entry in entries if entry.filename.endswith(".rels")]
        parser = etree.XMLParser(resolve_entities=False, no_network=True)
        for entry in relationship_files:
            root = etree.fromstring(archive.read(entry), parser=parser)
            for relation in root:
                if relation.get("TargetMode") != "External":
                    continue
                relation_type = relation.get("Type", "")
                target = relation.get("Target", "")
                if not relation_type.endswith("/hyperlink") or not target.startswith("https://"):
                    raise CatalogImportError("XLSX contains an unsafe external relationship")


def _trim_headers(values: tuple[str | None, ...]) -> tuple[str | None, ...]:
    result = values
    while result and result[-1] is None:
        result = result[:-1]
    return result


def _sheet_for_headers(
    headers: dict[str, tuple[str | None, ...]], expected: tuple[str, ...]
) -> str:
    matches = [name for name, values in headers.items() if values == expected]
    signature = hashlib.sha256(_canonical_json(expected).encode()).hexdigest()
    if len(matches) != 1:
        raise CatalogImportError(f"Expected exactly one sheet with schema {signature}")
    return matches[0]


def _sheet_for_header_variants(
    headers: dict[str, tuple[str | None, ...]], variants: tuple[tuple[str, ...], ...]
) -> tuple[str, tuple[str, ...]]:
    matches = [
        (name, variant)
        for name, values in headers.items()
        for variant in variants
        if values == variant
    ]
    if len(matches) != 1:
        signatures = [
            hashlib.sha256(_canonical_json(item).encode()).hexdigest() for item in variants
        ]
        raise CatalogImportError(
            f"Expected exactly one sheet with an approved schema: {signatures}"
        )
    return matches[0]


def _schema_signature(headers: dict[str, tuple[str | None, ...]]) -> str:
    canonical = {name: list(values) for name, values in sorted(headers.items())}
    return hashlib.sha256(_canonical_json(canonical).encode()).hexdigest()


def _parse_csts(sheet: Any) -> tuple[list[dict[str, Any]], list[StagingRow], list[ImportIssue]]:
    records: list[dict[str, Any]] = []
    staging: list[StagingRow] = []
    issues: list[ImportIssue] = []
    seen: set[str] = set()
    for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        raw = _raw_row(CST_HEADERS, values)
        if not any(value is not None for value in raw.values()):
            continue
        if raw["CST-IBS/CBS"] is None:
            staging.append(StagingRow(sheet.title, row_number, _json_safe(raw)))
            continue
        code = _code(raw["CST-IBS/CBS"], 3)
        description = _text(raw["Descrição CST-IBS/CBS"])
        row_errors: list[ImportIssue] = []
        if code is None or description is None:
            row_errors.append(
                _issue(
                    "MISSING_REQUIRED_FIELD",
                    "CST and description are required",
                    sheet.title,
                    row_number,
                )
            )
        elif code in seen:
            row_errors.append(
                _issue(
                    "DUPLICATE_CST", f"Duplicate CST {code}", sheet.title, row_number, "CST-IBS/CBS"
                )
            )
        else:
            seen.add(code)
            indicators = {
                _normalized_header(key): _indicator(value, key, sheet.title, row_number, row_errors)
                for key, value in raw.items()
                if key not in CST_HEADERS[:2]
            }
            records.append({"code": code, "description": description, "indicators": indicators})
        staging.append(StagingRow(sheet.title, row_number, _json_safe(raw), tuple(row_errors)))
    if not records:
        issues.append(_issue("NO_CST_ROWS", "No valid CST records found", sheet.title))
    return records, staging, issues


def _parse_classifications(
    sheet: Any, valid_csts: set[str], epoch: datetime, headers: tuple[str, ...]
) -> tuple[list[dict[str, Any]], list[StagingRow], list[ImportIssue]]:
    records: list[dict[str, Any]] = []
    staging: list[StagingRow] = []
    issues: list[ImportIssue] = []
    seen: set[str] = set()
    for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        raw = _raw_row(headers, values)
        if not any(value is not None for value in raw.values()):
            continue
        row_errors: list[ImportIssue] = []
        cst = _code(raw["CST-IBS/CBS"], 3)
        code = _code(raw["cClassTrib"], 6)
        name = _text(raw["Nome cClassTrib"])
        description = _text(raw["Descrição cClassTrib"])
        if None in (cst, code, name, description):
            row_errors.append(
                _issue(
                    "MISSING_REQUIRED_FIELD",
                    "CST, cClassTrib, name and description are required",
                    sheet.title,
                    row_number,
                )
            )
        elif code in seen:
            row_errors.append(
                _issue(
                    "DUPLICATE_CCLASSTRIB",
                    f"Duplicate cClassTrib {code}",
                    sheet.title,
                    row_number,
                    "cClassTrib",
                )
            )
        elif cst not in valid_csts:
            row_errors.append(
                _issue(
                    "UNKNOWN_CST",
                    f"cClassTrib {code} references unknown CST {cst}",
                    sheet.title,
                    row_number,
                    "CST-IBS/CBS",
                )
            )
        else:
            assert (
                cst is not None
                and code is not None
                and name is not None
                and description is not None
            )
            seen.add(code)
            records.append(
                {
                    "code": code,
                    "cst": cst,
                    "cst_description": _text(raw["Descrição CST-IBS/CBS"]),
                    "name": name,
                    "description": description,
                    "valid_from": _date_value(
                        raw["dIniVig"], epoch, "dIniVig", sheet.title, row_number, row_errors
                    ),
                    "valid_to": _date_value(
                        raw["dFimVig"], epoch, "dFimVig", sheet.title, row_number, row_errors
                    ),
                    "updated_on": _date_value(
                        raw["DataAtualização"],
                        epoch,
                        "DataAtualização",
                        sheet.title,
                        row_number,
                        row_errors,
                    ),
                    "attributes": {
                        _normalized_header(key): _normalized_value(value, epoch)
                        for key, value in raw.items()
                        if key
                        not in {
                            "CST-IBS/CBS",
                            "Descrição CST-IBS/CBS",
                            "cClassTrib",
                            "Nome cClassTrib",
                            "Descrição cClassTrib",
                            "dIniVig",
                            "dFimVig",
                            "DataAtualização",
                        }
                    },
                }
            )
        staging.append(StagingRow(sheet.title, row_number, _json_safe(raw), tuple(row_errors)))
    if not records:
        issues.append(
            _issue("NO_CCLASSTRIB_ROWS", "No valid cClassTrib records found", sheet.title)
        )
    return records, staging, issues


def _raw_row(headers: tuple[str, ...], values: tuple[Any, ...]) -> dict[str, Any]:
    trimmed = tuple(values)
    while len(trimmed) > len(headers) and trimmed[-1] is None:
        trimmed = trimmed[:-1]
    if len(trimmed) > len(headers):
        raise CatalogImportError("Row contains data beyond the declared schema")
    padded = trimmed + (None,) * (len(headers) - len(trimmed))
    return dict(zip(headers, padded, strict=True))


def _text(value: Any) -> str | None:
    if value is None:
        return None
    normalized = "\n".join(
        part.strip() for part in str(value).replace("\r\n", "\n").split("\n")
    ).strip()
    return normalized or None


def _code(value: Any, width: int) -> str | None:
    if value is None:
        return None
    if isinstance(value, int):
        return f"{value:0{width}d}"
    text = _text(value)
    return text.zfill(width) if text and text.isdigit() and len(text) <= width else text


def _indicator(
    value: Any, field: str, sheet: str, row: int, errors: list[ImportIssue]
) -> int | None:
    if value in (0, 1):
        return int(value)
    errors.append(_issue("INVALID_INDICATOR", f"{field} must be 0 or 1", sheet, row, field))
    return None


def _date_value(
    value: Any, epoch: datetime, field: str, sheet: str, row: int, errors: list[ImportIssue]
) -> str | None:
    if value is None:
        return None
    try:
        if isinstance(value, datetime):
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
        if isinstance(value, (int, float, Decimal)):
            converted = from_excel(value, epoch)
            return (
                converted.date().isoformat()
                if isinstance(converted, datetime)
                else converted.isoformat()
            )
        return date.fromisoformat(str(value)).isoformat()
    except (TypeError, ValueError, OverflowError):
        errors.append(_issue("INVALID_DATE", f"Invalid date in {field}", sheet, row, field))
        return None


def _normalized_value(value: Any, epoch: datetime) -> Any:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, float):
        return format(Decimal(str(value)), "f")
    if isinstance(value, str):
        return _text(value)
    return value


def _json_safe(raw: dict[str, Any]) -> dict[str, Any]:
    return {key: _normalized_value(value, datetime(1899, 12, 30)) for key, value in raw.items()}


def _normalized_header(value: str) -> str:
    return value.replace(" ", "")


def _canonical_json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def _issue(
    code: str,
    message: str,
    sheet: str | None = None,
    row: int | None = None,
    field: str | None = None,
) -> ImportIssue:
    return ImportIssue(code, message, sheet, row, field)
