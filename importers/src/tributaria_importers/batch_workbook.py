"""Parsing for user-uploaded batch classification spreadsheets (Etapa 23).

Unlike `ibs_cbs_catalog.py`/`ncm_catalog.py`/`nbs_catalog.py`, the input here
is **not** a trusted, manually-verified official artifact - it is a file an
end user uploads through the batch consultation screen. It gets the same
zip-bomb/XXE/unsafe-path defenses as the official importers (duplicated
rather than shared, following this codebase's existing convention of
independent screens/importers over premature abstraction), plus two checks
that do not apply to trusted government artifacts: rejection of any
macro-carrying container (`.xlsm`/`.xlsb`/`.xls`, or an `.xlsx` that embeds
`xl/vbaProject.bin`) and tolerant, accent/case-insensitive header mapping
(a user's own column names, not a government-published schema).

Column recognition is deliberately permissive (ADR-0027, item 9 of the
Etapa 23 specification: "não exigir que todas as colunas estejam
presentes"): any column that does not match a known alias is preserved
verbatim in `raw` (so nothing is silently dropped) and reported as an
"unknown header" - never an error that blocks the whole file. Extra columns
become `product_attributes` for the tax engine (see
`tributaria_api.application.batch_classification`), giving a user a way to
supply the scope-defining facts a specific rule needs
(e.g. a column literally named `buyer.health_entity_status`).
"""

from __future__ import annotations

import csv
import io
import unicodedata
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

MAX_UNCOMPRESSED_BYTES = 50_000_000
MAX_ZIP_ENTRIES = 200
FORBIDDEN_EXTENSIONS = frozenset({".xlsm", ".xlsb", ".xls", ".xltm"})
MACRO_ENTRY_NAMES = frozenset({"xl/vbaproject.bin"})

_HEADER_ALIASES: dict[str, str] = {
    "codigointerno": "internal_code",
    "codigo": "internal_code",
    "codigodoproduto": "internal_code",
    "codigointernodoproduto": "internal_code",
    "descricao": "description",
    "descricaodoproduto": "description",
    "ncm": "ncm",
    "nbs": "nbs",
    "ncmounbs": "ncm_or_nbs",
    "ncmnbs": "ncm_or_nbs",
    "tipodeobjeto": "object_kind",
    "tipoobjeto": "object_kind",
    "objeto": "object_kind",
    "dataoperacao": "operation_date",
    "datadaoperacao": "operation_date",
    "data": "operation_date",
}

_OBJECT_KIND_ALIASES: dict[str, str] = {
    "bem": "GOOD",
    "mercadoria": "GOOD",
    "produto": "GOOD",
    "good": "GOOD",
    "servico": "SERVICE",
    "service": "SERVICE",
}


class BatchImportError(ValueError):
    pass


@dataclass(frozen=True, slots=True)
class ParsedBatchRow:
    row_number: int
    raw: dict[str, str]
    internal_code: str | None
    description: str | None
    ncm: str | None
    nbs: str | None
    object_kind: str | None
    operation_date: date | None
    date_error: bool
    extra_attributes: dict[str, str]


@dataclass(frozen=True, slots=True)
class ParsedBatchWorkbook:
    rows: tuple[ParsedBatchRow, ...]
    total_row_count: int
    truncated: bool
    unknown_headers: tuple[str, ...]
    recognized_headers: tuple[str, ...]


def parse_batch_workbook(
    content: bytes, filename: str, *, max_rows: int
) -> ParsedBatchWorkbook:
    suffix = Path(filename).suffix.lower()
    if suffix in FORBIDDEN_EXTENSIONS:
        raise BatchImportError(
            f"Formato de planilha não suportado por segurança: {suffix} "
            "(apenas .xlsx e .csv são aceitos; formatos com macro são rejeitados)"
        )
    if suffix == ".csv":
        return _parse_csv(content, max_rows=max_rows)
    if suffix == ".xlsx":
        return _parse_xlsx(content, max_rows=max_rows)
    raise BatchImportError(f"Extensão de arquivo não suportada: {suffix or '(nenhuma)'}")


def _parse_xlsx(content: bytes, *, max_rows: int) -> ParsedBatchWorkbook:
    _validate_xlsx_container(content)
    buffer = io.BytesIO(content)
    workbook = load_workbook(buffer, read_only=True, data_only=True, keep_links=False)
    try:
        sheet = workbook.worksheets[0]
        rows_iter = sheet.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration as exc:
            raise BatchImportError("Planilha vazia: nenhum cabeçalho encontrado") from exc
        headers = [str(cell).strip() if cell is not None else "" for cell in header_row]
        field_by_column, unknown_headers = _map_headers(headers)
        rows: list[ParsedBatchRow] = []
        total = 0
        for values in rows_iter:
            if all(value is None for value in values):
                continue
            total += 1
            if total > max_rows:
                continue
            rows.append(
                _build_row(
                    row_number=total,
                    headers=headers,
                    values=[_cell_to_text(value) for value in values],
                    field_by_column=field_by_column,
                )
            )
        return ParsedBatchWorkbook(
            rows=tuple(rows),
            total_row_count=total,
            truncated=total > max_rows,
            unknown_headers=tuple(unknown_headers),
            recognized_headers=tuple(sorted(set(field_by_column.values()))),
        )
    finally:
        workbook.close()


def _parse_csv(content: bytes, *, max_rows: int) -> ParsedBatchWorkbook:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("cp1252")
    reader = csv.reader(io.StringIO(text), delimiter=_sniff_delimiter(text))
    try:
        headers = next(reader)
    except StopIteration as exc:
        raise BatchImportError("Planilha vazia: nenhum cabeçalho encontrado") from exc
    headers = [header.strip() for header in headers]
    field_by_column, unknown_headers = _map_headers(headers)
    rows: list[ParsedBatchRow] = []
    total = 0
    for values in reader:
        if not any(value.strip() for value in values):
            continue
        total += 1
        if total > max_rows:
            continue
        padded = values + [""] * (len(headers) - len(values))
        rows.append(
            _build_row(
                row_number=total,
                headers=headers,
                values=[value.strip() for value in padded[: len(headers)]],
                field_by_column=field_by_column,
            )
        )
    return ParsedBatchWorkbook(
        rows=tuple(rows),
        total_row_count=total,
        truncated=total > max_rows,
        unknown_headers=tuple(unknown_headers),
        recognized_headers=tuple(sorted(set(field_by_column.values()))),
    )


def _sniff_delimiter(text: str) -> str:
    first_line = text.splitlines()[0] if text.splitlines() else ""
    return ";" if first_line.count(";") > first_line.count(",") else ","


def _validate_xlsx_container(content: bytes) -> None:
    if not zipfile.is_zipfile(io.BytesIO(content)):
        raise BatchImportError("Arquivo não é um contêiner XLSX válido")
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        entries = archive.infolist()
        if len(entries) > MAX_ZIP_ENTRIES:
            raise BatchImportError("XLSX contém entradas em excesso")
        total_uncompressed = sum(entry.file_size for entry in entries)
        if total_uncompressed > MAX_UNCOMPRESSED_BYTES:
            raise BatchImportError("XLSX excede o limite de tamanho descomprimido")
        for entry in entries:
            name = entry.filename
            if ".." in Path(name).parts:
                raise BatchImportError("XLSX contém um caminho de arquivo inseguro")
            if name.lower() in MACRO_ENTRY_NAMES:
                raise BatchImportError(
                    "XLSX contém macro (VBA) - apenas planilhas sem macro são aceitas"
                )


def _normalize_header(text: str) -> str:
    decomposed = unicodedata.normalize("NFD", text)
    without_accents = "".join(ch for ch in decomposed if not unicodedata.combining(ch))
    return "".join(ch for ch in without_accents.casefold() if ch.isalnum())


def _map_headers(headers: list[str]) -> tuple[dict[int, str], list[str]]:
    field_by_column: dict[int, str] = {}
    unknown: list[str] = []
    for index, header in enumerate(headers):
        normalized = _normalize_header(header)
        field = _HEADER_ALIASES.get(normalized)
        if field is not None:
            field_by_column[index] = field
        elif header:
            unknown.append(header)
    return field_by_column, unknown


def _cell_to_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _build_row(
    *,
    row_number: int,
    headers: list[str],
    values: list[str],
    field_by_column: dict[int, str],
) -> ParsedBatchRow:
    raw = {headers[index]: values[index] for index in range(len(headers)) if headers[index]}
    fields: dict[str, str] = {}
    for index, field in field_by_column.items():
        if index < len(values) and values[index]:
            fields[field] = values[index]
    extra_attributes = {
        headers[index]: values[index]
        for index in range(len(headers))
        if headers[index] and index not in field_by_column and index < len(values) and values[index]
    }

    ncm = fields.get("ncm")
    nbs = fields.get("nbs")
    if "ncm_or_nbs" in fields:
        candidate = fields["ncm_or_nbs"]
        if candidate.replace(".", "").isdigit() and "." not in candidate:
            ncm = ncm or candidate
        else:
            nbs = nbs or candidate

    object_kind = None
    if "object_kind" in fields:
        object_kind = _OBJECT_KIND_ALIASES.get(_normalize_header(fields["object_kind"]), "OTHER")

    operation_date: date | None = None
    date_error = False
    if "operation_date" in fields:
        operation_date = _parse_date(fields["operation_date"])
        date_error = operation_date is None

    return ParsedBatchRow(
        row_number=row_number,
        raw=raw,
        internal_code=fields.get("internal_code"),
        description=fields.get("description"),
        ncm=ncm,
        nbs=nbs,
        object_kind=object_kind,
        operation_date=operation_date,
        date_error=date_error,
        extra_attributes=extra_attributes,
    )


def extra_attributes_from_raw(raw: dict[str, str]) -> dict[str, str]:
    """Recompute the non-recognized columns of an already-persisted raw row.

    Used when reprocessing a stored `classification_batch_rows.raw_values`
    (Etapa 23) instead of a freshly-parsed `ParsedBatchRow`, so both paths
    agree on exactly one definition of "known column" (`_HEADER_ALIASES`).
    """
    return {
        header: value
        for header, value in raw.items()
        if value and _normalize_header(header) not in _HEADER_ALIASES
    }


def _parse_date(value: str) -> date | None:
    for fmt in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            continue
    try:
        return date.fromisoformat(value)
    except ValueError:
        return None
