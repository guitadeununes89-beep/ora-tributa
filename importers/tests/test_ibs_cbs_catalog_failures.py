from pathlib import Path

import pytest
from openpyxl import Workbook
from tributaria_importers.ibs_cbs_catalog import (
    CCLASSTRIB_HEADERS,
    CST_HEADERS,
    CatalogImportError,
    import_official_catalog,
)


def synthetic_catalog(path: Path, *, duplicate: bool = False, missing_name: bool = False) -> None:
    workbook = Workbook()
    cst_sheet = workbook.active
    cst_sheet.title = "CST fixture"
    cst_sheet.append(CST_HEADERS)
    cst_sheet.append(["000", "Synthetic CST", *([0] * (len(CST_HEADERS) - 2))])
    class_sheet = workbook.create_sheet("cClass fixture")
    class_sheet.append(CCLASSTRIB_HEADERS)
    row: list[object | None] = [None] * len(CCLASSTRIB_HEADERS)
    row[0] = "000"
    row[1] = "Synthetic CST"
    row[2] = "000001"
    row[3] = None if missing_name else "Synthetic classification"
    row[4] = "Fixture only; not a legal rule"
    row[21] = "2040-01-01"
    row[23] = "2040-01-01"
    class_sheet.append(row)
    if duplicate:
        class_sheet.append(row)
    workbook.save(path)


def test_duplicate_and_missing_required_field_fail_closed(tmp_path: Path) -> None:
    duplicate = tmp_path / "duplicate.xlsx"
    synthetic_catalog(duplicate, duplicate=True)
    duplicate_result = import_official_catalog(duplicate)
    assert duplicate_result.valid is False
    assert any(
        error.code == "DUPLICATE_CCLASSTRIB"
        for row in duplicate_result.staging_rows
        for error in row.errors
    )

    missing = tmp_path / "missing.xlsx"
    synthetic_catalog(missing, missing_name=True)
    missing_result = import_official_catalog(missing)
    assert missing_result.valid is False
    assert any(
        error.code == "MISSING_REQUIRED_FIELD"
        for row in missing_result.staging_rows
        for error in row.errors
    )


def test_incompatible_official_schema_is_rejected(tmp_path: Path) -> None:
    artifact = tmp_path / "incompatible.xlsx"
    workbook = Workbook()
    workbook.active.append(["unknown", "schema"])
    workbook.save(artifact)

    with pytest.raises(CatalogImportError, match="Expected exactly one sheet"):
        import_official_catalog(artifact)
